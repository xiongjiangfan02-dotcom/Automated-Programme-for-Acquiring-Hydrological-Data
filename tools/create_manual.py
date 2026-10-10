from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "贵州毕节大型水库实时水情采集模块使用说明.docx"
BLACK = "000000"
NAVY = "1F4E78"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F4F6F7"
GRID = "D9D9D9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), GRID)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=None, color=BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, after=5, before=0, line=1.35, keep_next=False):
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(after)
    fmt.space_before = Pt(before)
    fmt.line_spacing = line
    fmt.keep_with_next = keep_next
    return paragraph


def add_text(doc, text, bold_lead=None):
    p = set_paragraph(doc.add_paragraph())
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        set_paragraph(p, after=3, line=1.25)
        set_run_font(p.add_run(item))


def add_numbered(doc, items):
    for number, item in enumerate(items, 1):
        p = doc.add_paragraph()
        set_paragraph(p, after=4, line=1.25)
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        set_run_font(p.add_run(f"{number}.  {item}"))


def add_code(doc, code):
    for line in code.splitlines():
        p = doc.add_paragraph()
        set_paragraph(p, after=0, line=1.05)
        p.paragraph_format.left_indent = Cm(0.6)
        r = p.add_run(line or " ")
        set_run_font(r, name="Consolas", size=9)
        p._p.get_or_add_pPr().append(OxmlElement("w:keepNext")) if line != code.splitlines()[-1] else None


def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(13 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    for run in p.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True)
    return p


def add_table(doc, headers, rows, widths=None, font_size=9):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(p, after=0, line=1.05)
        set_run_font(p.add_run(str(value)), size=font_size, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[i])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) <= 4 else WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph(p, after=0, line=1.08)
            set_run_font(p.add_run(str(value)), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, size=9, color="666666")


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(header.add_run("贵州毕节大型水库实时水情采集模块使用说明"), size=8.5, color="666666")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("第 "), size=9, color="666666")
    add_page_number(footer)
    set_run_font(footer.add_run(" 页"), size=9, color="666666")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(66)
    title.paragraph_format.space_after = Pt(20)
    set_run_font(title.add_run("贵州毕节大型水库实时水情采集模块使用说明"), size=20, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph(subtitle, after=36, line=1.2)
    set_run_font(subtitle.add_run("部署 运行 审计 故障处理和验收"), size=13, color="444444")

    meta_rows = [
        ("模块版本", "0.1.0"),
        ("文档日期", "2026年10月10日"),
        ("项目目录", r"D:\生态\bijie_hydro_collector"),
        ("数据范围", "毕节市大型水库实时水情"),
        ("当前结论", "模块可运行并通过28项测试；官网当前快照未发现毕节匹配记录，尚未通过有目标数据的生产验收"),
    ]
    add_table(doc, ["项目", "内容"], meta_rows, widths=[3.8, 11.5], font_size=9.5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("适用于获得官网自动采集和数据存储书面授权的运行人员及维护人员"), size=10, color="555555")
    doc.add_page_break()

    heading(doc, "文档导航", 1)
    add_bullets(doc, [
        "项目范围和当前状态",
        "目录结构和运行环境",
        "配置文件和白名单维护",
        "首次部署和模板验证",
        "单次采集和结果判断",
        "Excel 输出和批次归档",
        "每日定时任务配置",
        "质量检查和故障处理",
        "上传接口边界",
        "官网实测记录",
        "上线验收清单和日常检查表",
    ])
    add_text(doc, "本手册按实际操作顺序组织。首次部署人员应从运行环境开始阅读；已完成部署的值班人员可直接查看单次采集、结果判断和故障处理章节。")

    heading(doc, "1 项目范围和当前状态", 1)
    heading(doc, "1 1 采集范围", 2)
    add_text(doc, "本模块只处理水利部全国水雨情信息网的大型水库实时水情页面。MVP 输出对象是毕节市大型水库当前快照，不是毕节全部水文站，也不包含普通河道水位、流量、雨量或全部中小型水库数据。")
    add_text(doc, "来源页面为 http://xxfb.mwr.cn/sq_dxsk.html。程序先取得页面，再从页面脚本中发现当前公开数据接口。接口路径可能变化，因此默认配置不固定接口地址。程序只访问与来源页面相同的域名，拒绝跨域接口和跨域重定向。")

    heading(doc, "1 2 当前可用状态", 2)
    add_table(doc, ["项目", "状态", "说明"], [
        ("离线功能", "通过", "28项自动化测试通过，覆盖模板、字体、筛选、去重、归档、并发锁和失败路径"),
        ("官网页面", "间歇可访问", "2026年10月10日18时40分取得HTTP 200；18时45分复测连接被拒绝"),
        ("全国数据", "已验证", "成功批次取得571条记录，字段结构和类型检查通过"),
        ("贵州数据", "已验证", "当次响应含4条贵州记录，均完成字体解码"),
        ("毕节目标数据", "未发现", "当次响应未匹配夹岩、洪家渡、倒天河、象鼻岭候选"),
        ("正式成果", "未生成", "没有目标记录时不生成空Excel，不更新成功状态"),
        ("Seafile上传", "未启用", "仅保留接口边界，没有网络连接或上传行为"),
    ], widths=[3.1, 2.6, 9.6], font_size=9)
    add_text(doc, "当前版本可以用于受控试运行和审计验证。生产验收仍需等官网返回至少一条经官方名录核对的毕节大型水库记录，并由人工逐项比较库名、行政区划、观测时间、水位和日变幅。")

    heading(doc, "1 3 Excel 模板约束", 2)
    add_text(doc, r"模板路径为 D:\生态\获取\格式.xlsx。程序每次运行都以只读方式验证模板，不会修改原文件。2026年10月10日复核的SHA-256为 7694D0CD838575638FEC98A75F7A65AF7B5BE1CDF7C444B40E51A283DF6F7160。")
    add_table(doc, ["顺序", "字段", "输出规则"], [
        ("1", "流域", "官网 bsnm 解码值"),
        ("2", "行政区划", "官网 addvnm 解码值"),
        ("3", "河名", "官网 rvnm 解码值"),
        ("4", "库名", "白名单标准名称"),
        ("5", "时间", "北京时间，显示格式 yyyy-mm-dd hh:mm"),
        ("6", "库水位(米)", "数值，最多显示四位小数"),
        ("7", "日变幅(米)", "保留正负号，最多显示四位小数"),
    ], widths=[1.5, 4.2, 9.6])
    add_text(doc, "模板必须只包含 Sheet1、七列表头和一行表头，且不能增加隐藏列、合并单元格或公式。成果文件可以新增原始数据和采集日志工作表，但 Sheet1 的七列结构不得改变。")

    heading(doc, "2 目录结构和运行环境", 1)
    heading(doc, "2 1 项目目录", 2)
    add_table(doc, ["路径", "用途"], [
        (r"config\default.json", "正式采集配置"),
        (r"config\reservoir_allowlist.json", "经证据约束的毕节候选水库白名单"),
        (r"src\bijie_hydro", "采集、解码、筛选、质量检查、Excel和状态管理模块"),
        (r"tests", "自动化测试"),
        (r"references\official_reservoir_roster_2026.xlsx", "公开官方名录证据副本"),
        (r"requirements.lock.txt", "本机验证过的依赖版本"),
        (r"setup.ps1", "创建独立Python环境并安装依赖"),
        (r"run_once.ps1", "执行一次正式采集"),
        (r"docs", "使用说明文档"),
    ], widths=[6.0, 9.3])

    heading(doc, "2 2 数据目录", 2)
    add_table(doc, ["路径", "用途", "写入条件"], [
        (r"D:\生态\获取\格式.xlsx", "Excel模板", "只读，任何运行都不得写入"),
        (r"D:\生态\获取\成果\报表日期", "正式成果", "仅成功且非重复批次写入"),
        (r"D:\生态\获取\归档\批次日期\批次编号", "原始响应和审计证据", "每个批次均建立目录"),
        (r"D:\生态\获取\归档\fonts", "官网动态字体缓存", "取得新字体版本时写入"),
        (r"D:\生态\获取\归档\state", "最后成功快照状态", "成果成功保存后原子更新"),
    ], widths=[6.2, 5.0, 4.1], font_size=8.7)

    heading(doc, "2 3 运行环境", 2)
    add_bullets(doc, [
        "Windows 操作系统和 PowerShell。",
        "Python 3.11 或更高版本。当前项目独立环境使用 Python 3.12。",
        "fonttools 4.66.1、openpyxl 3.1.5、et-xmlfile 2.0.0、tzdata 2026.5。",
        "能够读取模板和项目目录，并能够写入成果及归档目录的 Windows 用户权限。",
        "能够访问目标官网的网络环境。无需也不应配置验证码绕过、代理池或登录规避。",
    ])

    heading(doc, "3 配置文件和白名单维护", 1)
    heading(doc, "3 1 正式配置", 2)
    add_table(doc, ["配置节", "主要参数", "运行要求"], [
        ("source", "page_url timeout_seconds max_attempts retry_base_seconds", "默认15秒、最多2次，页面和接口必须同源"),
        ("paths", "template output_root archive_root font_cache", "使用绝对路径，模板与成果不能相同"),
        ("font_decoder", "reference_fonts required_fields fail_on_unresolved", "目标相关字段不得关闭失败阻断"),
        ("filter", "province_names city_names allowlist_file", "白名单必须有已核实证据"),
        ("quality", "记录数量、时效和变化阈值", "只能根据试运行证据调整"),
        ("state", "file skip_duplicate_snapshot", "状态损坏时停止，禁止静默清空"),
        ("output", "文件名、日期和数值格式、工作表名称", "日期及数值格式不得偏离约定"),
        ("upload", "enabled provider", "本阶段 enabled 必须为 false"),
    ], widths=[3.2, 6.8, 5.3], font_size=8.5)

    heading(doc, "3 2 白名单规则", 2)
    add_text(doc, "白名单的用途是识别官网大型水库页面中的毕节记录，不是把名录中的所有水库强行输出。2026年公开名录同时包含大型和防汛重点中型水库，因此名录只用于名称和属地证据，最终大型分类仍以官网大型水库页面是否返回该记录为准。")
    add_table(doc, ["标准名称", "官方名录名称", "名录位置", "启用状态"], [
        ("夹岩水库", "夹岩水利枢纽", "第1383行 毕节市", "启用"),
        ("洪家渡水库", "洪家渡", "第1400行 毕节市", "启用"),
        ("倒天河水库", "倒天河", "第1401行 毕节市七星关区", "启用"),
        ("象鼻岭水库", "象鼻岭", "第1418行 毕节市", "启用"),
        ("附廓水库", "附廓", "第1402行 黔西市", "禁用 待复核"),
    ], widths=[3.4, 3.8, 5.0, 3.1], font_size=8.7)
    add_text(doc, "维护白名单时必须填写 canonical_name、aliases、province、city、evidence_status 和 source_url。启用项的 evidence_status 必须为已核实。站号 idNo 只能从官网实际响应中取得并经人工核对，不能按行政区划代码或记录顺序猜测。")
    add_text(doc, "未绑定站号时，程序使用受控库名和行政区划完整值匹配。匹配不使用任意子串，因此不会因为其他地区名称中包含毕节二字而误判。行政区划为空时，只允许已绑定且名称一致的站号匹配。")

    heading(doc, "4 首次部署和模板验证", 1)
    heading(doc, "4 1 创建独立环境", 2)
    add_text(doc, "在项目目录打开 PowerShell，然后执行以下命令。安装过程只修改项目下的 .venv，不修改系统 Python。")
    add_code(doc, 'Set-Location "D:\\生态\\bijie_hydro_collector"\n.\\setup.ps1')
    add_text(doc, "安装完成后检查依赖一致性。")
    add_code(doc, '.\\.venv\\Scripts\\python.exe -m pip check')
    add_text(doc, "正常结果为 No broken requirements found。")

    heading(doc, "4 2 验证模板", 2)
    add_code(doc, '.\\.venv\\Scripts\\python.exe -m bijie_hydro.cli --config config\\default.json validate-template')
    add_text(doc, "正常输出应显示 Sheet1、七个约定表头、max_row 为 1、max_column 为 7、hidden_columns 为空、formats 为 General、formulas 为空。任何差异都应先检查模板来源，不要直接修改程序放宽校验。")

    heading(doc, "4 3 运行测试", 2)
    add_code(doc, '.\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v')
    add_text(doc, "2026年10月10日当前版本运行 28 项测试，全部通过。测试数据存放在临时目录，不写入正式成果和历史状态。测试通过只能证明离线逻辑符合预期，不能代替真实官网和目标水库数据的人工验收。")

    heading(doc, "5 单次采集和结果判断", 1)
    heading(doc, "5 1 执行命令", 2)
    add_code(doc, 'Set-Location "D:\\生态\\bijie_hydro_collector"\n.\\run_once.ps1')
    add_text(doc, "脚本默认使用项目 .venv。也可以通过 -Config 指定其他配置，通过 -Python 指定解释器。正式运行不要把离线回放配置误用于生产状态目录。")

    heading(doc, "5 2 运行流程", 2)
    add_numbered(doc, [
        "取得单实例锁，防止两个采集任务同时写状态或成果。",
        "只读验证 Excel 模板，并将验证摘要写入批次归档。",
        "读取并验证白名单证据状态。",
        "访问大型水库页面，从页面发现唯一公开接口。",
        "请求接口并保存原始 JSON。",
        "下载或读取对应动态字体，按字形轮廓解码字段。",
        "按毕节白名单筛选，清洗时间和数值，统一处理缺测。",
        "执行字段、数量、时效、重复、历史差异和结构检查。",
        "成功时从模板生成临时 Excel，重读校验后改名为正式成果。",
        "原子更新最后成功状态，并写入批次日志和文件清单。",
    ])

    heading(doc, "5 3 批次状态", 2)
    add_table(doc, ["状态", "含义", "是否生成成果", "值班处理"], [
        ("成功", "取得目标记录并通过全部阻断检查", "是", "抽查成果和采集日志"),
        ("重复跳过", "数据与最后成功快照完全一致", "否", "正常保留批次审计"),
        ("无匹配数据", "官网成功返回，但没有毕节白名单匹配项", "否", "核对官网当前是否确无目标记录"),
        ("待复核", "历史变化、时间回退或阈值超限", "否或保留已完成文件", "人工核验后决定是否接受"),
        ("失败", "网络、结构、目标字体、模板或本地文件异常", "否", "按错误分类处理"),
    ], widths=[2.4, 6.1, 2.6, 4.2], font_size=8.6)

    heading(doc, "5 4 退出码", 2)
    add_table(doc, ["退出码", "含义", "自动重试建议"], [
        ("0", "成功或重复跳过", "不重试"),
        ("2", "采集、质量或配置失败", "先分类；普通瞬时网络故障最多补跑两次"),
        ("3", "命令入口遇到本地文件或网络系统错误", "检查路径、权限和磁盘"),
        ("4", "来源正常但没有匹配数据", "不立即重试，人工查看官网范围"),
        ("5", "401、403或429访问限制", "禁止短间隔自动重试"),
    ], widths=[2.2, 8.0, 5.1], font_size=8.8)

    heading(doc, "6 数据处理规则", 1)
    heading(doc, "6 1 时间", 2)
    add_bullets(doc, [
        "官网返回完整年份时直接使用该年份。",
        "仅当官网时间为 MM-DD HH:MM 时，先使用采集年份；若推定时间明显晚于采集时间超过30天，则回退一年。",
        "所有时间按 Asia Shanghai 时区处理，Excel 中写入不带时区标识的北京时间。",
        "官网观测时间写入 Sheet1，程序采集时间写入原始数据和采集日志，两者不得混用。",
        "成果目录和文件名中的报表日期使用采集日，不把一次实时快照解释为全天日报。",
    ])

    heading(doc, "6 2 数值和缺测", 2)
    add_table(doc, ["官网值", "主表值", "审计记录"], [
        ("有效非零数字", "按数值写入，四位小数四舍五入显示", "保留原始字符串和解码值"),
        ("负日变幅", "保留负号", "保留原始值"),
        ("-- 或空值", "空单元格", "记录缺测标记或空值"),
        ("任何数值零", "空单元格", "记录接口零值按缺测处理"),
        ("乱码或非数值", "不输出", "数值有效性检查失败并停止"),
        ("NaN 或无穷大", "不输出", "非有限数值检查失败并停止"),
    ], widths=[4.0, 5.6, 5.7], font_size=8.8)

    heading(doc, "6 3 动态字体", 2)
    add_text(doc, "官网使用动态字体标签编码中文和数字。程序下载本批次字体，将字形轮廓与本机参考字体比较，再恢复真实字符。字体 ID 变化不会复用旧映射；每个实际使用的字体文件会保存到批次归档。")
    add_text(doc, "全角标点等可直接识别的 Unicode 标点不依赖字形映射。非目标省份的未解析生僻字记录为警告，不阻止毕节结果。行政区划无法解析、贵州记录无法完整解析、已绑定站号或受控库名记录存在未解析字段时，系统停止输出。")

    heading(doc, "6 4 去重", 2)
    add_bullets(doc, [
        "同一站号和同一观测时间的完全相同记录可以合并。",
        "同一站号和同一观测时间出现不同字段值时视为冲突，停止输出。",
        "同一水库一次响应中有多个观测时间时，只保留最新观测。",
        "同一水库对应多个站号时停止并要求人工复核。",
        "整个目标快照与最后成功快照一致时，状态为重复跳过，不重复生成成果。",
    ])

    heading(doc, "7 Excel 输出和批次归档", 1)
    heading(doc, "7 1 成果命名", 2)
    add_code(doc, r'D:\生态\获取\成果\YYYYMMDD\毕节市大型水库实时水情_YYYYMMDD_批次编号.xlsx')
    add_text(doc, "程序先在目标目录写入隐藏前缀的临时文件，保存后重新打开检查 Sheet1 的列数和行数，再改名为最终文件。最终文件已存在时拒绝覆盖。写入中断不会留下冒充正式成果的半成品。")

    heading(doc, "7 2 工作表", 2)
    add_table(doc, ["工作表", "用途", "主要内容"], [
        ("Sheet1", "业务成果", "原模板七列，只写毕节匹配记录"),
        ("原始数据", "逐条审计", "全国原始记录、解码值、筛选状态、缺测原因和原始JSON摘要"),
        ("采集日志", "批次概览", "批次号、时间、接口、请求次数、记录数、质量检查和异常"),
    ], widths=[3.0, 3.8, 8.5], font_size=8.8)
    add_text(doc, "原始字符串强制以文本写入，避免以等号开头的内容被 Excel 当作公式。超过 Excel 单元格限制的审计内容会标记截断，完整内容仍在 raw_response.json。原始归档不做这种截断。")

    heading(doc, "7 3 批次归档", 2)
    add_table(doc, ["文件", "说明"], [
        ("source_page.html", "成功取得的来源页面"),
        ("raw_response.json", "接口原始响应，是回放和审计的首要证据"),
        ("font_字体标识.ttf", "本批次实际使用的动态字体"),
        ("decoded_records.json", "字段解码结果和逐条错误"),
        ("parsed_records.json", "筛选和缺测标记后的原始记录对象"),
        ("selected_records.json", "进入目标质量检查的标准记录"),
        ("reservoir_allowlist.json", "本批次实际使用的白名单快照"),
        ("template_validation.json", "模板结构检查结果"),
        ("provenance.json", "采集模式、时区、处理版本和模板校验和"),
        ("quality_checks.json", "结构、数量、时效、重复和历史差异检查"),
        ("batch_log.json", "批次最终状态和异常摘要"),
        ("manifest.json", "归档文件大小及SHA-256，检查后续内容是否变化"),
    ], widths=[5.2, 10.1], font_size=8.6)
    add_text(doc, "manifest.json 不是数字签名，不能证明文件由特定主体签发，也不能代替目录权限、备份和审计制度。归档目录应限制写权限，并纳入单位现有备份策略。")

    heading(doc, "8 每日定时任务配置", 1)
    heading(doc, "8 1 建议计划", 2)
    add_table(doc, ["项目", "建议值"], [
        ("执行频率", "每日一次"),
        ("执行时间", "北京时间09时10分"),
        ("运行账户", "有项目读取、成果和归档写入权限的固定Windows账户"),
        ("程序", r"C:\Windows\System32\WindowsPowerShell\v1.0\powershell.exe"),
        ("参数", r'-NoProfile -File "D:\生态\bijie_hydro_collector\run_once.ps1"'),
        ("起始目录", r"D:\生态\bijie_hydro_collector"),
        ("并发", "计划程序设置不启动新实例；模块内部另有文件锁"),
        ("普通网络故障", "10分钟后补跑，最多两次，由退出码分类控制"),
        ("访问限制", "退出码5，不进行短间隔补跑"),
    ], widths=[4.3, 11.0], font_size=8.6)

    heading(doc, "8 2 建立任务", 2)
    add_numbered(doc, [
        "打开 Windows 任务计划程序并选择创建任务。",
        "常规页使用固定运行账户，选择不管用户是否登录都要运行，并采用最低必要权限。",
        "触发器页建立每天09时10分触发器。",
        "操作页填写本节给出的程序、参数和起始目录。",
        "条件页根据服务器策略决定是否要求交流电源和网络可用。",
        "设置页选择如果任务正在运行则不启动新实例。不要设置所有失败都自动重试。",
        "创建后手动运行一次，检查退出码、最新批次目录和成果目录。",
        "将任务历史记录接入单位现有监控，按退出码区分通知级别。",
    ])
    add_text(doc, "模块不会自动创建计划任务，不会修改系统执行策略，也不会发送邮件或消息。通知应由现有任务监控、日志平台或告警系统读取退出码和 batch_log.json 后实现。")

    heading(doc, "9 质量检查和故障处理", 1)
    heading(doc, "9 1 主要检查", 2)
    add_table(doc, ["检查", "正常条件", "异常处理"], [
        ("页面和接口", "HTTP成功、业务码为0、result为数组", "失败并保存可获得的证据"),
        ("全国记录数", "100至2000条", "视为结构或数据异常并停止"),
        ("字段结构", "八个必需字段均存在且类型允许", "停止，检查官网变化"),
        ("动态字体", "目标相关记录完整解码", "停止，保存字体并人工分析"),
        ("目标记录", "白名单匹配且字段可清洗", "零条为无匹配数据，不生成空成果"),
        ("时间", "可解析、不过旧、不过度超前", "失败或待复核"),
        ("日变幅", "绝对值不超过配置阈值", "待复核"),
        ("历史水位", "与上次值差不超过阈值", "待复核"),
        ("历史时间", "观测时间不回退", "待复核"),
        ("成果结构", "Sheet1七列且行数一致", "不发布临时文件"),
    ], widths=[3.3, 5.5, 6.5], font_size=8.3)

    heading(doc, "9 2 常见故障", 2)
    add_table(doc, ["现象", "可能原因", "处理步骤"], [
        ("连接被拒绝或超时", "官网临时不可达、出口网络或防火墙", "保留失败批次；浏览器核对；普通故障稍后有限补跑；不要提高频率"),
        ("401或403", "访问权限、来源策略或安全设备", "停止自动重试；确认书面授权、出口地址和官网策略"),
        ("429", "请求频率受限", "记录Retry-After；停止短间隔重试；由管理员确定下次运行"),
        ("接口候选不唯一", "页面脚本结构变化", "查看source_page.html；人工确认公开接口后更新解析逻辑"),
        ("目标字体解码失败", "官网字体变化或参考字体无对应字形", "保存字体；核对原始字符；增加经验证参考字体或调整解码器并补测试"),
        ("全国非目标字形警告", "生僻字不在参考字体", "确认不涉及贵州或已绑定目标；保留警告，不手工猜字"),
        ("模板结构变化", "模板被改名、加列、加表或加入合并单元格", "恢复经批准模板或明确变更需求后修改程序"),
        ("历史状态损坏", "文件被截断或手工修改", "停止运行；从备份和最后成功批次恢复，禁止直接删除后继续"),
        ("磁盘写入失败", "空间、权限、杀毒锁定或路径问题", "检查磁盘和权限；确认无临时文件冒充成果后重跑"),
    ], widths=[4.0, 4.6, 6.7], font_size=8.0)

    heading(doc, "9 3 离线回放", 2)
    add_text(doc, "离线回放用于复现已归档的接口响应，不能代替新的官网采集。回放配置必须使用独立的 output_root、archive_root、state.file 和 font_cache，避免污染正式状态。带动态字体的响应需要对应字体已经缓存。")
    add_code(doc, '.\\.venv\\Scripts\\python.exe -m bijie_hydro.cli --config config\\replay.json collect --input-json "D:\\路径\\raw_response.json"')
    add_text(doc, "回放时 attempts 为 0，http_status 可以为空，这是离线模式的正常表现。provenance.json 会把模式记录为 offline_replay。")

    heading(doc, "10 上传接口边界", 1)
    add_text(doc, "当前采集流程不调用上传模块。upload.enabled 必须为 false，DisabledUploadPort 的所有方法都会拒绝执行。这样可以确保采集、归档和质量检查不依赖 Seafile 或其他云平台。")
    add_table(doc, ["接口", "用途", "建议返回值"], [
        ("upload_file(file_path metadata)", "上传已通过质量检查的成果文件", "上传任务ID"),
        ("upload_records(records metadata)", "可选的结构化记录上传", "上传任务ID"),
        ("check_upload_status(task_id)", "查询异步上传结果", "待处理 成功 失败"),
    ], widths=[5.4, 6.1, 3.8], font_size=8.6)
    add_text(doc, "后续上传任务应单独读取成功成果，使用目标资料库、目标目录和文件 SHA-256 组成幂等键。上传失败不删除成果，也不把成功采集改为失败。Seafile 配置预计包括服务器地址、访问凭证、资料库 ID、目录、文件名、标签和权限信息。凭证必须放入安全凭证存储，不能写入普通配置、Excel 或日志。")

    heading(doc, "11 官网实测记录", 1)
    heading(doc, "11 1 在线成功取数批次", 2)
    add_table(doc, ["项目", "结果"], [
        ("批次编号", "20261010T184057_cd8bc3c4"),
        ("时间", "2026年10月10日18时40分57秒至18时42分37秒 北京时间"),
        ("页面", "http://xxfb.mwr.cn/sq_dxsk.html"),
        ("发现接口", "http://xxfb.mwr.cn/OTMbxwsvgKjswqj/OTMkzpvsbSmpk"),
        ("HTTP状态", "200"),
        ("全国记录", "571条"),
        ("结构和类型", "通过"),
        ("原始结果", "旧阻断规则因17条全国字形异常停止，没有生成成果"),
    ], widths=[4.1, 11.2], font_size=8.8)
    add_text(doc, "分析显示，17条异常中16条属于被动态字体标签包裹的全角括号或逗号，1条属于非贵州记录的生僻字。贵州的4条记录均完整解码：乌江渡水电站、红枫水库、百花水库和天生桥一级。该响应没有夹岩、洪家渡、倒天河或象鼻岭记录。")

    heading(doc, "11 2 修正和真实响应回放", 2)
    add_text(doc, "程序随后允许可识别的 Unicode 标点直接通过，并把字体失败阻断范围收紧到贵州、已绑定站号、受控库名及行政区划未知的记录。非目标省份生僻字保留警告，不再阻断毕节结果。相关测试增加后，总测试数为28项并全部通过。")
    add_table(doc, ["项目", "结果"], [
        ("回放批次", "20261010T184431_29189cfe"),
        ("输入", "在线成功批次的原始JSON和真实字体缓存"),
        ("全国记录", "571条"),
        ("字体检查", "1条非目标记录警告，目标相关失败为0"),
        ("目标记录", "0条"),
        ("批次状态", "无匹配数据"),
        ("成果文件", "未生成"),
    ], widths=[4.1, 11.2], font_size=8.8)
    add_text(doc, "这次回放证明程序能够处理当天真实响应并正确区分无匹配数据。它不证明官网始终无毕节记录，也不证明所有毕节大型水库都应出现在当前页面。")

    heading(doc, "11 3 最后一次在线复测", 2)
    add_table(doc, ["项目", "结果"], [
        ("批次编号", "20261010T184547_aceb4219"),
        ("时间", "2026年10月10日18时45分47秒至18时46分24秒 北京时间"),
        ("请求次数", "2次"),
        ("结果", "WinError 10061 目标计算机积极拒绝连接"),
        ("批次状态", "失败"),
        ("正式成果", "未生成"),
    ], widths=[4.1, 11.2], font_size=8.8)
    add_text(doc, "因此，当前最新在线状态是连接失败；当前最新可用官网快照仍是18时40分成功取得的571条响应。值班人员不能把最后一次连接失败写成无数据，也不能把成功快照中没有毕节匹配项写成采集失败。")

    heading(doc, "12 上线验收清单", 1)
    add_table(doc, ["序号", "验收项", "合格条件", "当前状态"], [
        ("1", "模板保护", "模板校验和不变且结构检查通过", "通过"),
        ("2", "自动化测试", "全部测试通过", "28项通过"),
        ("3", "官网页面和接口", "连续多个采集日可访问并保留原始证据", "未完成"),
        ("4", "真实动态字体", "目标记录全部正确解码", "贵州记录已验证 无毕节样本"),
        ("5", "目标名录", "官网目标名称和idNo逐项核实", "名称已核实 idNo待绑定"),
        ("6", "目标数据人工比对", "至少一批毕节记录与页面逐字段一致", "未完成"),
        ("7", "Excel成果", "三张工作表、格式、缺测和正负号正确", "合成数据通过 真实目标待验"),
        ("8", "历史差异", "连续采集能识别重复、回退和异常变化", "测试通过 真实目标待验"),
        ("9", "计划任务", "固定账户、单实例和分类重试生效", "未创建"),
        ("10", "备份和告警", "归档纳入备份，退出码接入现有监控", "待部署"),
        ("11", "上传", "另行评审后启用", "本阶段不验收"),
    ], widths=[1.3, 3.8, 6.3, 3.9], font_size=8.0)
    add_text(doc, "生产上线前至少连续人工核验三个采集日。若三个采集日均无毕节记录，只能确认官网当前页面未返回目标，不能用空数据完成目标字段验收。此时应由业务负责人确认该页面是否确实覆盖目标大型水库，或是否需要另行取得经授权的数据源。")

    heading(doc, "13 日常运行检查表", 1)
    add_table(doc, ["检查时点", "检查内容", "通过标准"], [
        ("运行前", "模板存在、磁盘空间、项目环境、网络策略", "无缺失和阻断"),
        ("运行后", "退出码和最新 batch_log.json", "状态含义与任务结果一致"),
        ("成功批次", "Sheet1记录、时间、水位、日变幅、原始数据和日志", "抽样与官网一致"),
        ("无匹配数据", "全国记录数、贵州记录和白名单版本", "来源正常且确无匹配"),
        ("失败批次", "错误类别、请求次数和已有归档", "已分类并采取有限处理"),
        ("每日", "成果目录和归档目录", "没有半成品或异常重复"),
        ("每周", "磁盘、备份、失败趋势和官网结构变化", "容量正常且问题已跟踪"),
        ("白名单变更后", "证据、别名、行政区划和idNo", "双人复核并重新测试"),
    ], widths=[3.0, 7.2, 5.1], font_size=8.5)

    heading(doc, "14 合规和安全边界", 1)
    add_bullets(doc, [
        "仅采集公开页面、公开接口或书面授权范围内的数据。",
        "不绕过验证码、登录、访问控制、频率限制或网站安全措施。",
        "401、403和429视为停止信号，不使用代理轮换或伪装绕过。",
        "保留数据来源URL、采集时间、原始响应、处理版本和批次文件清单。",
        "采集结果可能受到官网缺测、延迟、结构变化和动态字体变化影响，不能承诺绝对准确。",
        "系统通过目录权限、最小权限和凭证隔离降低泄露风险，但不能承诺绝对不会泄露数据。",
        "公开来源不自动意味着可以无限传播。对外共享应符合书面授权和单位数据管理制度。",
        "生产上线前必须人工抽样核验，异常批次不能以普通成功结果发布。",
    ])

    heading(doc, "15 参考来源", 1)
    add_table(doc, ["来源", "地址或位置", "用途"], [
        ("全国水雨情信息网", "http://xxfb.mwr.cn/sq_dxsk.html", "大型水库实时水情来源页面"),
        ("2026年全国大型及防汛重点中型水库防汛责任人名单", "https://www.mem.gov.cn/gk/zfxxgkpt/fdzdgknr/202604/W020260426472645632108.xlsx", "水库名称和属地证据"),
        ("本地Excel模板", r"D:\生态\获取\格式.xlsx", "成果主表结构和字段顺序"),
        ("项目README", r"D:\生态\bijie_hydro_collector\README.md", "运行限制和最新实现说明"),
        ("批次归档", r"D:\生态\获取\归档\20261010", "官网实测原始证据和日志"),
    ], widths=[5.0, 7.3, 3.0], font_size=7.5)
    trailing = doc.paragraphs[-1]
    trailing.paragraph_format.space_before = Pt(0)
    trailing.paragraph_format.space_after = Pt(0)
    trailing.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    trailing.paragraph_format.line_spacing = Pt(1)
    set_run_font(trailing.add_run(" "), size=1, color="FFFFFF")

    doc.core_properties.title = "贵州毕节大型水库实时水情采集模块使用说明"
    doc.core_properties.subject = "采集模块部署运行审计故障处理和验收"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "贵州 毕节 大型水库 水情 数据采集"
    doc.core_properties.comments = "根据2026年10月10日项目实现和官网实测编制"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
