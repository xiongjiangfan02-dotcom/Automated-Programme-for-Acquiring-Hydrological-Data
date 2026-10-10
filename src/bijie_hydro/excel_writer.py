from __future__ import annotations

import json
import os
import re
import tempfile
from dataclasses import asdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from .errors import TemplateError
from .models import BatchLog, RawRecord, ReservoirRecord


EXPECTED_HEADERS = ["流域", "行政区划", "河名", "库名", "时间", "库水位(米)", "日变幅(米)"]
RAW_HEADERS = [
    "批次编号", "来源URL", "采集时间", "原始序号", "原始idNo", "原始流域", "原始行政区划", "原始河名",
    "原始库名", "原始时间", "原始库水位", "原始日变幅", "解码值", "解码错误", "筛选状态", "缺测原因", "原始JSON",
]
LOG_HEADERS = [
    "批次编号", "任务状态", "开始时间", "结束时间", "来源页面", "接口URL", "HTTP状态", "请求次数", "全国记录数",
    "毕节记录数", "输出记录数", "报表日期", "质量检查", "异常摘要", "处理版本",
]


def _append(sheet, values):
    safe = []
    for value in values:
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, default=str)
        if isinstance(value, str):
            value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)
            if len(value) > 32767:
                value = value[:32690] + " [内容过长，完整值见批次原始JSON归档]"
        safe.append(value)
    sheet.append(safe)


def validate_template(path: Path) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise TemplateError("缺少openpyxl依赖") from exc
    if not path.is_file():
        raise TemplateError(f"模板文件不存在: {path}")
    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise TemplateError("模板缺少Sheet1")
        sheet = workbook["Sheet1"]
        headers = [sheet.cell(1, column).value for column in range(1, 8)]
        if headers != EXPECTED_HEADERS:
            raise TemplateError(f"Sheet1表头不一致: {headers!r}")
        if sheet.max_column != 7 or sheet.max_row != 1 or sheet.merged_cells or len(workbook.sheetnames) != 1:
            raise TemplateError("模板结构发生变化：要求仅Sheet1、七列、一行表头、无合并单元格")
        return {
            "sheets": workbook.sheetnames, "headers": headers, "max_row": sheet.max_row, "max_column": sheet.max_column,
            "hidden_columns": [key for key, value in sheet.column_dimensions.items() if value.hidden],
            "formats": [cell.number_format for cell in sheet[1]],
            "formulas": [cell.coordinate for row in sheet for cell in row if cell.data_type == "f"],
        }
    finally:
        workbook.close()


def _style_header(sheet, width_by_column: dict[int, float]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column, width in width_by_column.items():
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _populate_workbook(
    template: Path,
    destination: Path,
    records: list[ReservoirRecord],
    raw_records: list[RawRecord],
    log: BatchLog,
    source_url: str,
    collected_at,
    date_format: str,
    number_format: str,
    raw_sheet_name: str,
    log_sheet_name: str,
) -> Path:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment

    workbook = load_workbook(template)
    sheet = workbook["Sheet1"]
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    for record in records:
        _append(sheet, [
            record.basin,
            record.administrative_area,
            record.river_name,
            record.reservoir_name,
            record.observation_time.replace(tzinfo=None),
            float(record.water_level.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)) if record.water_level is not None else None,
            float(record.daily_change.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)) if record.daily_change is not None else None,
        ])
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 5).number_format = date_format
        sheet.cell(row, 6).number_format = number_format
        sheet.cell(row, 7).number_format = number_format
        for column in range(1, 8):
            sheet.cell(row, column).alignment = Alignment(vertical="center")

    for name in (raw_sheet_name, log_sheet_name):
        if name in workbook.sheetnames:
            del workbook[name]

    raw_sheet = workbook.create_sheet(raw_sheet_name)
    raw_sheet.append(RAW_HEADERS)
    for item in raw_records:
        _append(raw_sheet, [
            log.batch_id,
            source_url,
            collected_at.replace(tzinfo=None),
            item.index,
            item.values.get("idNo"),
            item.values.get("bsnm"),
            item.values.get("addvnm"),
            item.values.get("rvnm"),
            item.values.get("stnm"),
            item.values.get("tm"),
            item.values.get("rz"),
            item.values.get("rzRange"),
            json.dumps(item.decoded, ensure_ascii=False, default=str),
            "; ".join(item.decode_errors),
            item.filter_status,
            "; ".join(item.missing_reasons),
            json.dumps(item.values, ensure_ascii=False, default=str),
        ])
    raw_sheet.column_dimensions["M"].width = 40
    raw_sheet.column_dimensions["Q"].width = 60
    _style_header(raw_sheet, {1: 24, 2: 42, 3: 20, 4: 10, 5: 16, 6: 16, 7: 16, 8: 16, 9: 20, 10: 20, 11: 16, 12: 16, 14: 30, 15: 18, 16: 30})

    log_sheet = workbook.create_sheet(log_sheet_name)
    log_sheet.append(LOG_HEADERS)
    _append(log_sheet, [
        log.batch_id,
        log.status,
        log.started_at.replace(tzinfo=None),
        log.finished_at.replace(tzinfo=None) if log.finished_at else None,
        log.source_page,
        log.endpoint_url,
        log.http_status,
        log.attempts,
        log.national_count,
        log.target_count,
        log.output_count,
        log.report_date,
        json.dumps([asdict(check) for check in log.checks], ensure_ascii=False),
        "; ".join(log.errors),
        log.version,
    ])
    _style_header(log_sheet, {1: 24, 2: 14, 3: 20, 4: 20, 5: 42, 6: 48, 7: 12, 8: 12, 9: 14, 10: 14, 11: 14, 12: 14, 13: 60, 14: 50, 15: 14})
    for audit_sheet in (raw_sheet, log_sheet):
        for row in audit_sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        audit_sheet.sheet_view.showGridLines = False

    # Source text is literal text, including strings beginning with '='.
    for ws in workbook:
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        if ws.title != "Sheet1":
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if hasattr(cell.value, "year"):
                        cell.number_format = date_format

    workbook.save(destination)
    workbook.close()
    return destination


def write_workbook(template: Path, destination: Path, records, raw_records, log, source_url,
                   collected_at, date_format: str, number_format: str, raw_sheet_name: str,
                   log_sheet_name: str) -> Path:
    validate_template(template)
    if destination.resolve() == template.resolve():
        raise TemplateError("成果路径不得覆盖原模板")
    if destination.exists():
        raise TemplateError("成果文件已存在，禁止覆盖")
    if {raw_sheet_name, log_sheet_name} != {"原始数据", "采集日志"}:
        raise TemplateError("审计工作表名称错误")
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary_name = tempfile.mkstemp(prefix=".pending_", suffix=".xlsx", dir=destination.parent)
    os.close(fd)
    temporary = Path(temporary_name)
    try:
        _populate_workbook(template, temporary, records, raw_records, log, source_url,
                           collected_at, date_format, number_format, raw_sheet_name, log_sheet_name)
        from openpyxl import load_workbook
        verification = load_workbook(temporary, read_only=True)
        try:
            if verification["Sheet1"].max_column != 7 or verification["Sheet1"].max_row != len(records) + 1:
                raise TemplateError("保存后的成果结构校验失败")
        finally:
            verification.close()
        # On Windows rename refuses to overwrite; collection_lock prevents competing collectors.
        temporary.rename(destination)
    finally:
        temporary.unlink(missing_ok=True)
    return destination

