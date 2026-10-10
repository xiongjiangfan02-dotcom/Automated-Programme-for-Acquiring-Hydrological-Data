from __future__ import annotations

import copy
import hashlib
import io
import json
import tempfile
import unittest
from datetime import datetime, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock, patch
from urllib.error import HTTPError
from zoneinfo import ZoneInfo

from fontTools.fontBuilder import FontBuilder
from fontTools.pens.ttGlyphPen import TTGlyphPen
from openpyxl import Workbook, load_workbook

from bijie_hydro.config import load_settings
from bijie_hydro.errors import ConfigurationError, DecodeError, QualityError, SourceError, TemplateError
from bijie_hydro.excel_writer import EXPECTED_HEADERS, validate_template, write_workbook
from bijie_hydro.filtering import ReservoirAllowlist
from bijie_hydro.font_decoder import DynamicFontDecoder
from bijie_hydro.http_client import HttpClient, discover_endpoint, validate_source_url
from bijie_hydro.locking import collection_lock
from bijie_hydro.models import BatchLog, ReservoirRecord
from bijie_hydro.normalize import parse_observation_time
from bijie_hydro.pipeline import CollectionPipeline
from bijie_hydro.state import StateStore, snapshot_fingerprint
from bijie_hydro.quality import check_historical_time


TZ = ZoneInfo("Asia/Shanghai")
PROJECT = Path(__file__).resolve().parents[1]


def fixture(root):
    template = root / "template.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.active.append(EXPECTED_HEADERS)
    wb.save(template)
    wb.close()
    config = json.loads((PROJECT / "config/default.json").read_text(encoding="utf-8"))
    config["paths"] = {"template": str(template), "output_root": str(root / "outputs"),
                       "archive_root": str(root / "archive"), "font_cache": str(root / "fonts")}
    config["source"].update(page_url="http://offline.invalid/sq_dxsk.html", timeout_seconds=1, max_attempts=1)
    config["state"]["file"] = str(root / "state.json")
    config["filter"]["allowlist_file"] = str(root / "allowlist.json")
    config["quality"]["min_national_records"] = 1
    config["font_decoder"]["reference_fonts"] = []
    allowlist = {"version": "test", "reservoirs": [{"canonical_name": "夹岩水库", "aliases": ["夹岩"],
        "id_nos": ["id-1"], "province": "贵州省", "city": "毕节市", "evidence_status": "已核实",
        "source_url": "https://official.example/roster"}]}
    (root / "allowlist.json").write_text(json.dumps(allowlist, ensure_ascii=False), encoding="utf-8")
    config_path = root / "config.json"
    config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
    record = {"bsnm": "长江", "addvnm": "贵州省", "rvnm": "六冲河", "stnm": "夹岩",
              "tm": datetime.now(TZ).strftime("%Y-%m-%d %H:%M:%S"), "rz": "1000.12567", "rzRange": "-0.12", "idNo": "id-1"}
    return config_path, record


def response_file(root, records):
    path = root / "response.json"
    path.write_text(json.dumps({"returncode": 0, "result": records}, ensure_ascii=False), encoding="utf-8")
    return path


def synthetic_font(cmap, composite=False):
    """Synthetic fixtures; never used as water ministry data or decoding evidence."""
    order = [".notdef"] + list(dict.fromkeys(cmap.values()))
    glyphs = {}
    for index, name in enumerate(order):
        pen = TTGlyphPen(None)
        pen.moveTo((0, 0))
        pen.lineTo((100 + index * 25, 0))
        pen.lineTo((50, 100 + index * 10))
        pen.closePath()
        glyphs[name] = pen.glyph()
    if composite:
        base = order[1]
        order.append("component_base")
        glyphs["component_base"] = glyphs[base]
        pen = TTGlyphPen(glyphs)
        pen.addComponent("component_base", (1, 0, 0, 1, 0, 0))
        glyphs[base] = pen.glyph()
    builder = FontBuilder(1000, isTTF=True)
    builder.setupGlyphOrder(order)
    builder.setupCharacterMap(cmap)
    builder.setupGlyf(glyphs)
    builder.setupHorizontalMetrics({name: (500, 0) for name in order})
    builder.setupHorizontalHeader(ascent=800, descent=-200)
    builder.setupNameTable({"familyName": "Synthetic", "styleName": "Regular"})
    builder.setupOS2(sTypoAscender=800, sTypoDescender=-200, usWinAscent=800, usWinDescent=200)
    builder.setupPost()
    buffer = io.BytesIO()
    builder.save(buffer)
    return buffer.getvalue()


class FontTests(unittest.TestCase):
    def test_composite_and_rotating_font_ids(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            ref = root / "reference.ttf"
            ref.write_bytes(synthetic_font({ord("夹"): "a", ord("岩"): "b", ord("1"): "c"}))
            (root / "first.ttf").write_bytes(synthetic_font({0x3400: "a", 0x3401: "b", 0x3402: "c"}, composite=True))
            (root / "second.ttf").write_bytes(synthetic_font({0x3402: "a", 0x3401: "b", 0x3400: "c"}, composite=True))
            decoder = DynamicFontDecoder(Mock(), "http://offline.invalid", root, [ref], offline=True)
            self.assertEqual(decoder.decode("#firstotltag㐀㐁㐂#FontTag"), "夹岩1")
            self.assertEqual(decoder.decode("#secondotltag㐂㐁㐀#FontTag"), "夹岩1")
            with self.assertRaises(DecodeError):
                decoder.decode("#firstotltag㐃#FontTag")

    def test_failed_font_is_not_requested_for_every_record(self):
        with tempfile.TemporaryDirectory() as temp:
            client = Mock()
            client.request.side_effect = SourceError("blocked")
            decoder = DynamicFontDecoder(client, "http://offline.invalid", Path(temp), [])
            for _ in range(3):
                with self.assertRaises(DecodeError):
                    decoder.decode("#fontotltag㐀#FontTag")
            client.request.assert_called_once()

    def test_fullwidth_punctuation_passes_through(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            ref = root / "reference.ttf"
            ref.write_bytes(synthetic_font({ord("夹"): "a", ord("岩"): "b"}))
            (root / "font.ttf").write_bytes(synthetic_font({0x3400: "a", 0x3401: "b"}))
            decoder = DynamicFontDecoder(Mock(), "http://offline.invalid", root, [ref], offline=True)
            self.assertEqual(decoder.decode("#fontotltag㐀（㐁）#FontTag"), "夹（岩）")

    def test_encoded_pipeline_uses_cached_font_without_network(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config_path, record = fixture(root)
            ref = root / "reference.ttf"
            ref.write_bytes(synthetic_font({ord("夹"): "a", ord("岩"): "b", ord("1"): "c", ord("2"): "d", ord("3"): "e", ord("4"): "f", ord("-"): "g"}))
            cache = root / "fonts"
            cache.mkdir()
            (cache / "font.ttf").write_bytes(synthetic_font({0x3400: "a", 0x3401: "b", 0x3402: "c", 0x3403: "d", 0x3404: "e", 0x3405: "f", 0x3406: "g"}, composite=True))
            config = json.loads(config_path.read_text(encoding="utf-8"))
            config["font_decoder"]["reference_fonts"] = [str(ref)]
            config_path.write_text(json.dumps(config), encoding="utf-8")
            record["stnm"] = "#fontotltag㐀㐁#FontTag"
            record["rz"] = "#fontotltag㐂㐃.㐄㐅#FontTag"
            record["rzRange"] = "#fontotltag㐆0.㐂㐃#FontTag"
            pipeline = CollectionPipeline(load_settings(config_path))
            pipeline.client.request = Mock(side_effect=AssertionError("offline must not use network"))
            log = pipeline.run(response_file(root, [record]))
            self.assertEqual(log.status, "成功")
            self.assertEqual(len(list((root / "archive").rglob("font_font.ttf"))), 1)
            wb = load_workbook(log.output_file)
            self.assertEqual(wb["Sheet1"]["F2"].value, 12.34)
            self.assertEqual(wb["Sheet1"]["G2"].value, -0.12)
            wb.close()
            manifest = json.loads(next((root / "archive").rglob("manifest.json")).read_text(encoding="utf-8"))
            self.assertTrue(any(item["file"] == "raw_response.json" for item in manifest["files"]))


class PipelineFailureTests(unittest.TestCase):
    def test_non_target_decode_error_does_not_block_no_match_result(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config_path, record = fixture(root)
            record.update(addvnm="湖北", stnm="#missingotltag㐀#FontTag", idNo="other")
            config = json.loads(config_path.read_text(encoding="utf-8"))
            (root / "fonts").mkdir()
            (root / "fonts" / "missing.ttf").write_bytes(synthetic_font({0x3401: "dummy", 0x3400: "unknown"}))
            reference = root / "reference.ttf"
            reference.write_bytes(synthetic_font({ord("夹"): "known"}))
            config["font_decoder"]["reference_fonts"] = [str(reference)]
            config_path.write_text(json.dumps(config), encoding="utf-8")
            log = CollectionPipeline(load_settings(config_path)).run(response_file(root, [record]))
            self.assertEqual(log.status, "无匹配数据")
            check = next(item for item in log.checks if item.name == "动态字体解码")
            self.assertEqual(check.status, "警告")

    def test_conflicting_duplicate_preserves_error_without_workbook(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config, record = fixture(root)
            conflict = dict(record, rz="1001")
            with self.assertRaises(QualityError):
                CollectionPipeline(load_settings(config)).run(response_file(root, [record, conflict]))
            self.assertEqual(list((root / "outputs").rglob("*.xlsx")), [])
            logs = list((root / "archive").rglob("batch_log.json"))
            self.assertIn("冲突", logs[0].read_text(encoding="utf-8"))

    def test_malformed_numeric_is_not_silently_missing(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config, record = fixture(root)
            record["rz"] = "garbled"
            with self.assertRaises(QualityError):
                CollectionPipeline(load_settings(config)).run(response_file(root, [record]))
            self.assertFalse((root / "state.json").exists())

    def test_historical_jump_requires_review_and_retains_previous_state(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config, record = fixture(root)
            pipeline = CollectionPipeline(load_settings(config))
            pipeline.run(response_file(root, [record]))
            previous = (root / "state.json").read_bytes()
            record["rz"] = "1100"
            with self.assertRaises(QualityError):
                pipeline.run(response_file(root, [record]))
            self.assertEqual(previous, (root / "state.json").read_bytes())
            self.assertEqual(len(list((root / "outputs").rglob("*.xlsx"))), 1)
            self.assertTrue(any('"status": "待复核"' in p.read_text(encoding="utf-8") for p in (root / "archive").rglob("batch_log.json")))

    def test_empty_target_is_distinct_from_source_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config, record = fixture(root)
            record.update(stnm="其他水库", idNo="other", addvnm="浙江省")
            log = CollectionPipeline(load_settings(config)).run(response_file(root, [record]))
            self.assertEqual(log.status, "无匹配数据")
            self.assertEqual(log.national_count, 1)
            self.assertIsNone(log.output_file)

    def test_actual_http_attempt_count_is_preserved_on_failure(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config, _ = fixture(root)
            pipeline = CollectionPipeline(load_settings(config))
            pipeline.client.opener.open = Mock(side_effect=HTTPError("http://offline.invalid", 403, "Forbidden", {}, None))
            with self.assertRaises(SourceError):
                pipeline.run()
            log = json.loads(next((root / "archive").rglob("batch_log.json")).read_text(encoding="utf-8"))
            self.assertEqual(log["attempts"], 1)
            self.assertEqual(log["status"], "失败")


class SafetyTests(unittest.TestCase):
    def test_lock_releases_after_exception_and_rejects_parallel_run(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "lock"
            with collection_lock(path):
                with self.assertRaises(Exception):
                    with collection_lock(path):
                        self.fail("parallel collection entered")
            with collection_lock(path):
                pass

    def test_corrupt_state_fails_closed(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "state.json"
            path.write_text("broken", encoding="utf-8")
            with self.assertRaises(QualityError):
                StateStore(path).load()

    def test_unverified_allowlist_fails_closed(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            fixture(root)
            path = root / "allowlist.json"
            data = json.loads(path.read_text(encoding="utf-8"))
            data["reservoirs"][0]["evidence_status"] = "待核实"
            path.write_text(json.dumps(data), encoding="utf-8")
            with self.assertRaises(ConfigurationError):
                ReservoirAllowlist.load(path)

    def test_output_is_atomic_template_preserved_and_text_is_not_formula(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config, record = fixture(root)
            template = root / "template.xlsx"
            original = hashlib.sha256(template.read_bytes()).hexdigest()
            record["bsnm"] = "=1+1"
            log = CollectionPipeline(load_settings(config)).run(response_file(root, [record]))
            wb = load_workbook(log.output_file)
            self.assertEqual(wb["Sheet1"]["A2"].data_type, "s")
            self.assertEqual(wb["Sheet1"]["F2"].value, 1000.1257)
            wb.close()
            self.assertEqual(original, hashlib.sha256(template.read_bytes()).hexdigest())
            now = datetime.now(TZ)
            target = root / "failed.xlsx"
            rec = ReservoirRecord("长江", "贵州省", "六冲河", "夹岩水库", now, Decimal("123.4"), None, "id-1", 1)
            with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=OSError("disk full")):
                with self.assertRaises(OSError):
                    write_workbook(template, target, [rec], [], BatchLog("batch", "成功", now), "api", now,
                                   "yyyy-mm-dd hh:mm", "0.####", "原始数据", "采集日志")
            self.assertFalse(target.exists())
            self.assertEqual(list(root.glob(".pending_*")), [])

    def test_changed_template_columns_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            fixture(root)
            path = root / "template.xlsx"
            wb = load_workbook(path)
            wb["Sheet1"]["H1"] = "new"
            wb.save(path)
            wb.close()
            with self.assertRaises(TemplateError):
                validate_template(path)

    def test_short_leap_day_and_year_rollover(self):
        self.assertEqual(parse_observation_time("02-29 08:00", datetime(2024, 2, 29, 9, tzinfo=TZ)).day, 29)
        self.assertEqual(parse_observation_time("12-31 23:00", datetime(2027, 1, 1, 1, tzinfo=TZ)).year, 2026)

    def test_historical_time_rollback_is_review_required(self):
        now = datetime.now(TZ)
        record = ReservoirRecord("长江", "贵州省", "河名", "夹岩水库", now - timedelta(hours=1), None, None, "id", 1)
        result = check_historical_time([record], {"records": {"id": {"observation_time": now.isoformat()}}})
        self.assertEqual(result.status, "需复核")

    def test_snapshot_ignores_decimal_padding_but_tracks_source_metadata(self):
        from dataclasses import replace
        now = datetime.now(TZ)
        record = ReservoirRecord("长江", "贵州省", "河名", "夹岩水库", now, Decimal("123.4000"), None, "id", 1)
        self.assertEqual(snapshot_fingerprint([record]), snapshot_fingerprint([replace(record, water_level=Decimal("123.4"))]))
        self.assertNotEqual(snapshot_fingerprint([record]), snapshot_fingerprint([replace(record, river_name="修订河名")]))


class HttpSafetyTests(unittest.TestCase):
    def test_429_and_403_stop_without_retries(self):
        for code in (429, 403):
            client = HttpClient(1, 3, 0, "test")
            client.opener.open = Mock(side_effect=HTTPError("http://test.invalid", code, "Denied", {"Retry-After": "600"}, None))
            with self.assertRaises(SourceError):
                client.request("http://test.invalid")
            self.assertEqual(client.total_attempts, 1)

    def test_cross_origin_and_ambiguous_endpoint_are_rejected(self):
        with self.assertRaises(SourceError):
            validate_source_url("http://other.invalid/api", "http://source.invalid/page")
        with self.assertRaises(SourceError):
            discover_endpoint("http://source.invalid/page", "url:'/one',url:'/two'")


if __name__ == "__main__":
    unittest.main()
