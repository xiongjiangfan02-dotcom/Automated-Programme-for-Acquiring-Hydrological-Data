from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import load_workbook, Workbook

from bijie_hydro.config import load_settings
from bijie_hydro.pipeline import CollectionPipeline
from bijie_hydro.excel_writer import EXPECTED_HEADERS


class OfflinePipelineTests(unittest.TestCase):
    def test_plain_json_creates_result_and_audit_sheets(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            template = root / "template.xlsx"
            source = Workbook()
            source.active.title = "Sheet1"
            source.active.append(EXPECTED_HEADERS)
            source.save(template)
            source.close()
            response = root / "response.json"
            response.write_text(
                json.dumps(
                    {
                        "returncode": 0,
                        "message": "ok",
                        "result": [
                            {
                                "bsnm": "长江",
                                "addvnm": "贵州省毕节市",
                                "rvnm": "六冲河",
                                "stnm": "夹岩",
                                "tm": datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M:%S"),
                                "rz": "1131.2500",
                                "rzRange": 0,
                                "idNo": "id-1",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            config = root / "config.json"
            config.write_text(
                json.dumps(
                    {
                        "source": {
                            "page_url": "http://offline.invalid/sq_dxsk.html",
                            "endpoint_override": "http://offline.invalid/api",
                            "timeout_seconds": 1,
                            "max_attempts": 1,
                            "retry_base_seconds": 0,
                            "user_agent": "test",
                        },
                        "paths": {
                            "template": str(template),
                            "output_root": str(root / "outputs"),
                            "archive_root": str(root / "archive"),
                            "font_cache": str(root / "fonts"),
                        },
                        "font_decoder": {
                            "reference_fonts": [r"C:\Windows\Fonts\msyh.ttc"],
                            "required_fields": ["bsnm", "addvnm", "rvnm", "stnm", "idNo", "rz", "rzRange"],
                            "fail_on_unresolved": True,
                        },
                        "filter": {
                            "province_names": ["贵州省"],
                            "city_names": ["毕节市"],
                            "allowlist_file": str(root / "allowlist.json"),
                            "require_allowlist_match": True,
                        },
                        "quality": {
                            "min_national_records": 1,
                            "max_national_records": 10,
                            "max_observation_age_hours": 72,
                            "max_future_minutes": 30,
                            "max_level_change_meters": 20,
                            "max_water_level_delta_meters": 20,
                            "require_at_least_one_target": True,
                        },
                        "state": {"file": str(root / "state.json"), "skip_duplicate_snapshot": True},
                        "output": {
                            "file_name": "result_{report_date}_{batch_id}.xlsx",
                            "date_format": "yyyy-mm-dd hh:mm",
                            "number_format": "0.####",
                            "raw_sheet": "原始数据",
                            "log_sheet": "采集日志",
                        },
                        "upload": {"enabled": False, "provider": None},
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (root / "allowlist.json").write_text(
                json.dumps(
                    {
                        "version": "test",
                        "reservoirs": [
                            {
                                "canonical_name": "夹岩水库",
                                "aliases": ["夹岩"],
                                "id_nos": ["id-1"],
                                "province": "贵州省",
                                "city": "毕节市",
                                "evidence_status": "已核实",
                                "source_url": "https://official.example/roster",
                            }
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            log = CollectionPipeline(load_settings(config)).run(response)
            self.assertEqual(log.status, "成功")
            outputs = list((root / "outputs").rglob("*.xlsx"))
            self.assertEqual(len(outputs), 1)
            workbook = load_workbook(outputs[0], data_only=False)
            self.assertEqual(workbook.sheetnames, ["Sheet1", "原始数据", "采集日志"])
            self.assertEqual(workbook["Sheet1"].cell(2, 4).value, "夹岩水库")
            self.assertIsNone(workbook["Sheet1"].cell(2, 7).value)
            workbook.close()
            second = CollectionPipeline(load_settings(config)).run(response)
            self.assertEqual(second.status, "重复跳过")
            self.assertEqual(second.attempts, 0)
            self.assertEqual(len(list((root / "outputs").rglob("*.xlsx"))), 1)


if __name__ == "__main__":
    unittest.main()
