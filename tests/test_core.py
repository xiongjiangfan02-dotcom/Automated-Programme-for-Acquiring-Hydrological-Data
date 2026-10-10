from __future__ import annotations

import json
import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zoneinfo import ZoneInfo

from bijie_hydro.config import load_settings
from bijie_hydro.errors import ConfigurationError
from bijie_hydro.excel_writer import EXPECTED_HEADERS, validate_template, write_workbook
from bijie_hydro.filtering import matches_province, normalize_province_names
from bijie_hydro.http_client import discover_endpoint, parse_api_payload
from bijie_hydro.models import BatchLog, RawRecord, ReservoirRecord
from bijie_hydro.normalize import parse_measurement, parse_observation_time


TZ = ZoneInfo("Asia/Shanghai")


class HttpTests(unittest.TestCase):
    def test_endpoint_discovery_tracks_page_script(self):
        html = "<script>$.ajax({type:'POST',url:'/current/api',dataType:'json'})</script>"
        self.assertEqual(discover_endpoint("http://example.test/sq_dxsk.html", html), "http://example.test/current/api")

    def test_payload_validation(self):
        records, payload = parse_api_payload(json.dumps({"returncode": 0, "result": [{"stnm": "夹岩水库"}]}).encode())
        self.assertEqual(len(records), 1)
        self.assertEqual(payload["returncode"], 0)


class NormalizeTests(unittest.TestCase):
    def test_zero_and_markers_are_missing(self):
        for value in (0, 0.0, "0", "0.0000", "--", "", None):
            parsed, reason = parse_measurement(value)
            self.assertIsNone(parsed)
            self.assertTrue(reason)

    def test_signed_measurement_is_preserved(self):
        self.assertEqual(parse_measurement("-0.1250")[0], Decimal("-0.1250"))

    def test_full_and_short_time(self):
        collected = datetime(2026, 1, 1, 9, 0, tzinfo=TZ)
        self.assertEqual(parse_observation_time("2026-01-01 08:00:00", collected).year, 2026)
        self.assertEqual(parse_observation_time("12-31 08:00", collected).year, 2025)


class ProvinceFilterTests(unittest.TestCase):
    def test_exact_and_detailed_guizhou_areas_match(self):
        province = normalize_province_names(["贵州", "贵州省"])
        self.assertTrue(matches_province("贵州", province))
        self.assertTrue(matches_province("贵州省", province))
        self.assertTrue(matches_province("贵州省毕节市", province))
        self.assertFalse(matches_province("云南省", province))
        self.assertFalse(matches_province("云南省贵州村", province))
        self.assertFalse(matches_province("", province))


class ScheduleConfigTests(unittest.TestCase):
    def test_default_schedule_is_daily_at_nine(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = json.loads((Path(__file__).resolve().parents[1] / "config/default.json").read_text(encoding="utf-8"))
            config_path = root / "config.json"
            config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            settings = load_settings(config_path)
            self.assertEqual(settings.data["schedule"]["daily_time"], "09:00")
            self.assertEqual(settings.data["schedule"]["daily_count"], 1)

    def test_daily_count_limits_and_backward_compatibility(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = json.loads((Path(__file__).resolve().parents[1] / "config/default.json").read_text(encoding="utf-8"))
            config_path = root / "config.json"
            for count in (1, 2, 24):
                config["schedule"]["daily_count"] = count
                config_path.write_text(json.dumps(config), encoding="utf-8")
                self.assertEqual(load_settings(config_path).data["schedule"]["daily_count"], count)
            for count in (0, 25, True, False, 1.5, "2", None):
                with self.subTest(count=count):
                    config["schedule"]["daily_count"] = count
                    config_path.write_text(json.dumps(config), encoding="utf-8")
                    with self.assertRaises(ConfigurationError):
                        load_settings(config_path)
            del config["schedule"]["daily_count"]
            config_path.write_text(json.dumps(config), encoding="utf-8")
            load_settings(config_path)

    def test_invalid_schedule_time_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = json.loads((Path(__file__).resolve().parents[1] / "config/default.json").read_text(encoding="utf-8"))
            config["schedule"]["daily_time"] = "9:00"
            config_path = root / "config.json"
            config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(ConfigurationError):
                load_settings(config_path)

    def test_unsafe_task_name_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = json.loads((Path(__file__).resolve().parents[1] / "config/default.json").read_text(encoding="utf-8"))
            config["schedule"]["task_name"] = "bad/name"
            config_path = root / "config.json"
            config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(ConfigurationError):
                load_settings(config_path)

    def test_invalid_popup_settings_are_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            config = json.loads((Path(__file__).resolve().parents[1] / "config/default.json").read_text(encoding="utf-8"))
            config["schedule"]["popup_timeout_seconds"] = 0
            config_path = root / "config.json"
            config_path.write_text(json.dumps(config, ensure_ascii=False), encoding="utf-8")
            with self.assertRaises(ConfigurationError):
                load_settings(config_path)


class ExcelTests(unittest.TestCase):
    def test_template_and_output_structure(self):
        with tempfile.TemporaryDirectory() as temp:
            from openpyxl import Workbook
            template = Path(temp) / "template.xlsx"
            source = Workbook()
            source.active.title = "Sheet1"
            source.active.append(EXPECTED_HEADERS)
            source.save(template)
            source.close()
            info = validate_template(template)
            self.assertEqual(info["headers"], EXPECTED_HEADERS)
            output = Path(temp) / "result.xlsx"
            now = datetime(2026, 10, 10, 9, 0, tzinfo=TZ)
            log = BatchLog("batch", "成功", now, finished_at=now, source_page="page", endpoint_url="api", target_count=1, output_count=1)
            record = ReservoirRecord("长江", "贵州毕节", "六冲河", "洪家渡水库", now, Decimal("1131.2500"), None, "id-1", 1)
            raw = RawRecord(1, {"idNo": "raw", "rz": "0", "rzRange": "--"}, decoded={"idNo": "id-1"}, filter_status="命中", missing_reasons=["日变幅:缺测标记"])
            write_workbook(template, output, [record], [raw], log, "api", now, "yyyy-mm-dd hh:mm", "0.####", "原始数据", "采集日志")
            from openpyxl import load_workbook

            workbook = load_workbook(output, data_only=False)
            self.assertEqual(workbook.sheetnames, ["Sheet1", "原始数据", "采集日志"])
            sheet = workbook["Sheet1"]
            self.assertEqual([sheet.cell(1, col).value for col in range(1, 8)], EXPECTED_HEADERS)
            self.assertEqual(sheet.cell(2, 7).value, None)
            self.assertEqual(sheet.cell(2, 5).number_format, "yyyy-mm-dd hh:mm")
            self.assertEqual(sheet.cell(2, 6).number_format, "0.####")
            workbook.close()


if __name__ == "__main__":
    unittest.main()
